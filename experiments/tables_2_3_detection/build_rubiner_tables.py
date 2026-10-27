"""RubiNER_comparison.xlsx: RubiNER (SIESTA with its general NER fine-tuned on the
S-CECILIA training seeds, prediction file aad-fixC-ft-siesta-rep) against

  sheet "1 Same training data"   the D17 models fine-tuned on the same S-CECILIA seeds
  sheet "2 D17 configuration"    the D17 models as configured in D17.1 (transformers
                                 fine-tuned on CECILIA, spaCy as published), plus the
                                 D17 transformers as published

on S-CECILIA, CECILIA and CyNER, with P, R, exact R, F1, exact F1 and leak.
Bold: better than RubiNER on that measure (precision not marked).

Run from the repository root:  .venv/bin/python experiments/tables_2_3_detection/build_rubiner_tables.py
"""
import json, os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

import sys
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT)); import paths  # noqa: E402,F401  (puts every code folder on sys.path)
os.chdir(ROOT)
A = ROOT / "scores"
DS = [("S-CECILIA (held-out seeds)", "results_scec_test.json", "leak_private"),
      ("CECILIA (D17 test split)", "results_cecilia_test.json", "leak"),
      ("CyNER", "results_cyner_all.json", "leak")]
RES = {f: {r["system"]: r for r in json.load(open(f))} for _, f, _ in DS}
SC = {p.stem.removeprefix("scores_"): json.load(open(p)) for p in A.glob("scores_*.json")}
RUBI = "aad-fixC-ft-siesta-rep"


def rec(k, f):
    return SC[k].get(f) if k in SC else RES[f].get(k)


NAMES = {RUBI: "RubiNER (SIESTA, general NER fine-tuned on S-CECILIA)", "aad-fixC": "SIESTA without fine-tuning",
         "ft-cecilia__dslim__bert-base-NER": "dslim/bert-base-NER (D17 configuration)",
         "ft-cecilia__dbmdz__bert-large-cased-finetuned-conll03-english": "dbmdz/bert-large-cased-finetuned-conll03-english (D17 configuration)",
         "ft-cecilia__Jean-Baptiste__roberta-large-ner-english": "Jean-Baptiste/roberta-large-ner-english (D17 configuration)"}
FT_S = [f"{m} fine-tuned on S-CECILIA ({l})" for m in ("dslim/bert-base-NER", "dbmdz/bert-large-cased-finetuned-conll03-english",
        "Jean-Baptiste/roberta-large-ner-english") for l in ("conll4", "native")]
SHEETS = {
 "1 Same training data": [
   ("D17 models fine-tuned on the S-CECILIA training seeds (663 reports); conll4 = D17's 4 types, native = S-CECILIA's 31", FT_S),
   ("RubiNER and SIESTA", [RUBI, "aad-fixC"])],
 "2 D17 configuration": [
   ("D17 configuration as reported in D17.1: transformers fine-tuned on CECILIA (3 of 10 reproduced), spaCy as published",
    ["ft-cecilia__dslim__bert-base-NER", "ft-cecilia__dbmdz__bert-large-cased-finetuned-conll03-english",
     "ft-cecilia__Jean-Baptiste__roberta-large-ner-english", "spacy/en_core_web_sm", "spacy/en_core_web_md", "spacy/en_core_web_lg", "spacy/en_core_web_trf"]),
   ("D17 transformers as published (no fine-tuning)", ["Jean-Baptiste/roberta-large-ner-english", "FacebookAI/xlm-roberta-large-finetuned-conll03-english",
     "Davlan/bert-base-multilingual-cased-ner-hrl", "Davlan/xlm-roberta-base-ner-hrl", "Davlan/xlm-roberta-large-ner-hrl"]),
   ("RubiNER and SIESTA", [RUBI, "aad-fixC"])],
}
MET = [("P", "detect", "p"), ("R", "detect", "r"), ("Exact R", "exact", "r"), ("F1", "detect", "f1"), ("Exact F1", "exact", "f1"), ("Leak", None, None)]
F = "Arial"; thin = Side(style="thin", color="999999")
HDR = PatternFill("solid", fgColor="D9E1F2"); GRP = PatternFill("solid", fgColor="EDEDED"); RB = PatternFill("solid", fgColor="FFF2CC")


def val(r, m, lk):
    if r is None: return None
    return r.get(lk, r["leak"]) if m[1] is None else r[m[1]][m[2]]


wb = Workbook(); wb.remove(wb.active)
for title, groups in SHEETS.items():
    ws = wb.create_sheet(title)
    ws.cell(1, 1, {"1 Same training data": "RubiNER against the D17 models fine-tuned on the same S-CECILIA training data",
                   "2 D17 configuration": "RubiNER against the D17 models as configured and reported in D17.1"}[title]).font = Font(name=F, bold=True, size=12)
    ws.cell(3, 1, "Model"); ws.merge_cells(start_row=3, end_row=4, start_column=1, end_column=1)
    c = 2
    for lab, *_ in DS:
        ws.cell(3, c, lab); ws.merge_cells(start_row=3, end_row=3, start_column=c, end_column=c + 5)
        for i, (h, *_x) in enumerate(MET): ws.cell(4, c + i, h)
        c += 6
    last = c - 1; row = 5; rrows = []
    for g, systems in groups:
        ws.cell(row, 1, g).font = Font(name=F, bold=True, italic=True, size=10)
        for cc in range(1, last + 1): ws.cell(row, cc).fill = GRP
        row += 1
        for s in systems:
            ws.cell(row, 1, NAMES.get(s, s)).font = Font(name=F, size=10, bold=s == RUBI)
            c = 2
            for lab, f, lk in DS:
                r, ref = rec(s, f), rec(RUBI, f)
                for m in MET:
                    v = val(r, m, lk)
                    cell = ws.cell(row, c, "–" if v is None else round(v, 4 if m[0] == "Leak" else 3))
                    better = v is not None and s != RUBI and m[0] != "P" and ((v < val(ref, m, lk)) if m[0] == "Leak" else (v > val(ref, m, lk)))
                    cell.font = Font(name=F, size=10, bold=better or s == RUBI)
                    if v is not None: cell.number_format = "0.0%" if m[0] == "Leak" else "0.000"
                    c += 1
            if s == RUBI: rrows.append(row)
            row += 1
    for rr in ws.iter_rows(min_row=3, max_row=row - 1, max_col=last):
        for cell in rr:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if cell.row in (3, 4): cell.font = Font(name=F, size=10, bold=True); cell.fill = HDR; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif cell.column > 1: cell.alignment = Alignment(horizontal="center")
            if cell.row in rrows: cell.fill = RB
    notes = ["Bold: better than RubiNER on that measure (precision not marked). RubiNER row highlighted.",
             "P, R, F1: type-agnostic detection, 1:1 match at overlap (IoU) >= 0.5, micro-averaged. Exact R, Exact F1: identical boundaries. Leak: share of non-space gold-entity characters left uncovered; on S-CECILIA over personal and technical identifiers only.",
             "RubiNER = SIESTA with dslim/bert-base-NER replaced by the same model fine-tuned on the 37 S-CECILIA training seeds (SIESTA entity types as labels); cyber NER, entity pack, dates and patterns unchanged. Variant chosen on 4 validation seeds; evaluated in a copy of SIESTA, not yet in the repository.",
             "S-CECILIA test = 247 reports from 15 held-out seed incidents; none was used for training or any choice."]
    if title.startswith("2"):
        notes.append("On S-CECILIA RubiNER was trained on the same corpus (other seeds) while the D17 models were not, so S-CECILIA favours RubiNER in this comparison; CECILIA favours the D17 models, which were trained on it.")
    for k, t in enumerate(notes): ws.cell(row + 1 + k, 1, t).font = Font(name=F, size=9, italic=True)
    ws.column_dimensions["A"].width = 66
    for cc in range(2, last + 1): ws.column_dimensions[get_column_letter(cc)].width = 8.5
    ws.row_dimensions[3].height = 30; ws.freeze_panes = "B5"
wb.save(ROOT / "RubiNER_comparison.xlsx"); print("saved RubiNER_comparison.xlsx")
